import docx, re, sys, shutil
from docx.oxml.ns import qn
SRC=sys.argv[1]; DST=sys.argv[2]
d=docx.Document(SRC)
body=d.element.body
log=[]
# 1. hidden chars
n=0
for t in body.iter(qn('w:t')):
    if t.text and ('​' in t.text or '⁠' in t.text):
        n+=t.text.count('​')+t.text.count('⁠')
        t.text=t.text.replace('​','').replace('⁠','')
log.append(f"hidden chars removed: {n}")
# 2. minus in tables
n=0
for tbl in body.iter(qn('w:tbl')):
    for t in tbl.iter(qn('w:t')):
        if t.text:
            new=re.sub(r'(?<![\w])-(?=\d)','−',t.text)
            if new!=t.text: n+=1; t.text=new
for t in body.iter(qn('w:t')):
    if t.text=='-242': t.text='−242'; n+=1
log.append(f"minus signs fixed: {n}")

def para_replace(p, old, new):
    runs=p.runs
    full=''.join(r.text for r in runs)
    idx=full.find(old)
    if idx<0: return False
    s,e=idx,idx+len(old); pos=0; first=True
    for r in runs:
        t=r.text; rs,re_=pos,pos+len(t); pos=re_
        if re_<=s or rs>=e:
            if not(rs==re_==s and first): continue
        if first and rs<=s<=re_:
            tail=t[e-rs:] if e<=re_ else ''
            r.text=t[:s-rs]+new+tail; first=False
        elif not first and rs<e:
            r.text=t[e-rs:] if e<re_ else ''
    return True

def rep(old,new,count=1,scope=None):
    hits=0
    for i,p in enumerate(d.paragraphs):
        if scope and not scope(p): continue
        while old in ''.join(r.text for r in p.runs) and hits<(count*50 if count else 10**6):
            if para_replace(p,old,new): hits+=1
            else: break
            if old in new: break
    if hits!=count and count>0: log.append(f"!! {hits}/{count} :: {old[:70]}")
    return hits

R=[]
def r(old,new,count=1): R.append((old,new,count))

# ---- Tóm tắt / Mở đầu
r("Chuyên đề trình bày đầy đủ cơ sở lý thuyết cho từng kỹ thuật được sử dụng: CNN, MobileNetV2, Vision Transformer, RoPE, GQA, RMSNorm, các kỹ thuật XAI, và kiểm định thống kê có hiệu chỉnh Holm. Bên cạnh khảo sát các nghiên cứu liên quan, chuyên đề còn kiểm toán thiết kế theo từng thành phần: ablation huấn luyện lại từ đầu, kiểm định McNemar ghép cặp từng ảnh với hiệu chỉnh Holm, benchmark điều kiện suy giảm ở thời điểm kiểm thử, đánh giá thích nghi miền few-shot trên bộ dữ liệu độc lập, và ba kỹ thuật XAI thuộc ba nhóm cơ chế khác nhau. Mục đích là cho thấy từng lựa chọn thiết kế đã được kiểm chứng ra sao, thay vì chỉ báo cáo một mô hình đã lắp ráp hoàn chỉnh.",
  "Chuyên đề trình bày cơ sở lý thuyết cho từng kỹ thuật được sử dụng: CNN, MobileNetV2, Vision Transformer, RoPE, GQA, RMSNorm, các kỹ thuật XAI và kiểm định thống kê có hiệu chỉnh Holm. Mỗi thành phần của mô hình đều được kiểm tra riêng. Các biến thể được huấn luyện lại từ đầu sau khi bỏ từng thành phần, rồi so với mô hình đầy đủ bằng kiểm định McNemar trên từng ảnh. Chuyên đề cũng đo độ bền của mô hình khi ảnh bị nhòe hoặc thiếu sáng, đánh giá khả năng thích nghi trên một bộ dữ liệu độc lập, và dùng ba kỹ thuật XAI thuộc ba nhóm cơ chế để xem mô hình dựa vào vùng nào của ảnh.")
r("Học sâu, đặc biệt là mạng nơ-ron tích chập (CNN), đã chứng minh hiệu quả cao trong bài toán phân loại bệnh cây trồng từ ảnh. Các mạng tích chập nhẹ có ưu thế về số tham số và mang thiên kiến quy nạp cục bộ (local inductive bias) phù hợp với đặc trưng thị giác; một mạng đủ sâu vẫn có thể đạt trường nhận thức lớn, nhưng không trực tiếp mô hình hóa quan hệ theo cặp giữa mọi vị trí trong ảnh theo cách mà cơ chế tự chú ý (self-attention) làm được, điều quan trọng với các triệu chứng lan tỏa như xoăn lá hay vàng lá, nơi vùng bệnh không tập trung tại một vị trí cục bộ mà trải khắp bề mặt lá. Vision Transformer khai thác trực tiếp quan hệ giữa các token nhờ tự chú ý, nhưng các cấu hình tiêu chuẩn thường có chi phí tính toán cao hơn đáng kể so với CNN nhẹ. Kiến trúc lai CNN-Transformer vì vậy là một hướng dung hòa hai đặc tính này.",
  "Học sâu, đặc biệt là mạng nơ-ron tích chập (CNN), đã cho kết quả tốt trong bài toán phân loại bệnh cây trồng từ ảnh [27], [35]. Các mạng tích chập nhẹ có ít tham số và mang thiên kiến quy nạp cục bộ (local inductive bias), phù hợp với đặc trưng thị giác. Hạn chế của chúng là mỗi lớp chỉ tổng hợp thông tin trong một vùng nhỏ. Mạng sâu có thể mở rộng trường nhận thức, nhưng vẫn không mô hình hóa trực tiếp quan hệ giữa mọi cặp vị trí trong ảnh như cơ chế tự chú ý (self-attention). Điều này ảnh hưởng tới các bệnh như xoăn lá hay vàng lá, vì vùng bệnh trải khắp bề mặt lá. Vision Transformer xử lý tốt quan hệ toàn cục, song các cấu hình tiêu chuẩn tốn nhiều tính toán hơn CNN nhẹ. Kiến trúc lai CNN-Transformer là một cách dung hòa hai hướng này.")
r("Trong nhóm công trình được khảo sát ở Chương 2, các kết quả được báo cáo trên những bộ dữ liệu và quy trình khác nhau, nên chưa tạo thành một phép so sánh chung phân tích đồng thời đóng góp của từng thành phần kiến trúc, phản ứng trước suy giảm tổng hợp, và khoảng cách miền nền cảnh. Khoảng trống ở đây không phải là sự vắng mặt tuyệt đối của từng phân tích riêng lẻ, mà là sự vắng mặt của một quy trình kết hợp các trục đánh giá đó trên cùng một lần chia dữ liệu và cùng một quy trình huấn luyện. Đặc biệt, độ chính xác cao trên một bộ dữ liệu có nền được kiểm soát không đảm bảo hiệu năng trên ảnh có nền tự nhiên. Đây là khoảng trống mà đề tài này hướng tới giải quyết: xây dựng một mô hình lai CNN-Transformer đủ nhẹ để triển khai trên thiết bị biên, đồng thời kiểm toán một cách có hệ thống từng lựa chọn thiết kế bằng ablation huấn luyện lại từ đầu, kiểm định thống kê có hiệu chỉnh cho việc kiểm định nhiều lần, và đánh giá khả năng chuyển miền trên một bộ dữ liệu độc lập, thay vì chỉ báo cáo một mô hình đã lắp ráp hoàn chỉnh.",
  "Các công trình khảo sát ở Chương 2 báo cáo kết quả trên những bộ dữ liệu và quy trình khác nhau. Một số bài phân tích đóng góp của từng thành phần kiến trúc, một số đo độ bền trước ảnh suy giảm, một số xét khoảng cách miền. Điều còn thiếu là một quy trình gộp các phân tích đó trên cùng một lần chia dữ liệu và cùng một công thức huấn luyện. Ngoài ra, độ chính xác cao trên ảnh có nền được kiểm soát chưa đảm bảo mô hình chạy tốt trên ảnh nền tự nhiên. Đề tài hướng tới khoảng trống này theo hai hướng. Thứ nhất, xây dựng một mô hình lai CNN-Transformer đủ nhẹ để chạy trên thiết bị biên. Thứ hai, kiểm tra từng lựa chọn thiết kế bằng ablation huấn luyện lại từ đầu, kiểm định thống kê có hiệu chỉnh cho nhiều phép so sánh, và đánh giá chuyển miền trên một bộ dữ liệu độc lập.")
r("tập trung vào bài toán phân loại ảnh (classification), không mở rộng sang phát hiện đối tượng (detection) hay phân vùng (segmentation).","tập trung vào bài toán phân loại ảnh, không mở rộng sang phát hiện đối tượng hay phân vùng.")
# ---- Chương 1
r("Chương này trình bày nền tảng lý thuyết cho các kỹ thuật được sử dụng trong chuyên đề: bài toán phân loại ảnh bệnh cây trồng, mạng tích chập và kiến trúc MobileNetV2, Vision Transformer, ba kỹ thuật tiết kiệm tham số vay mượn từ LLM (RoPE, GQA, RMSNorm), cơ chế Triplet Attention và khối RES_Block của kiến trúc nền tảng, các kỹ thuật XAI, và các phương pháp kiểm định thống kê được dùng để so sánh mô hình ở Chương 4.",
  "Phần lý thuyết dưới đây đi theo thứ tự các thành phần của mô hình, từ xương sống MobileNetV2 tới khối Transformer, sau đó là các kỹ thuật XAI và các phép kiểm định thống kê dùng ở Chương 4.")
r("tách một phép tích chập chuẩn thành hai bước: depthwise convolution","một phép tích chập chuẩn được tách thành hai bước là depthwise convolution")
r("vì cơ chế tự chú ý (self-attention) vốn không","vì cơ chế tự chú ý vốn không")
r("Cơ chế tự chú ý tính toán, với mỗi token, một tổ hợp trọng số của mọi token khác trong chuỗi: mỗi token được chiếu thành ba vector: truy vấn (query), khóa (key) và giá trị (value); độ tương đồng",
  "Với mỗi token, cơ chế tự chú ý tính một tổ hợp có trọng số của mọi token khác trong chuỗi. Cụ thể, mỗi token được chiếu thành ba vector: truy vấn (query), khóa (key) và giá trị (value). Độ tương đồng")
r("Nhược điểm là độ phức tạp tính toán của tự chú ý tăng theo bậc hai của số token, và không có thiên kiến quy nạp cục bộ (inductive bias) như tích chập, nên thường cần nhiều dữ liệu huấn luyện hơn hoặc phải khởi tạo từ một xương sống CNN đã học đặc trưng cục bộ, chính là cách tiếp cận của ChiliEdgeNet.",
  "Nhược điểm là chi phí của tự chú ý tăng theo bình phương số token. Transformer cũng thiếu thiên kiến quy nạp cục bộ như tích chập, nên thường cần nhiều dữ liệu hơn [6], hoặc phải đặt sau một xương sống CNN đã học sẵn đặc trưng cục bộ. ChiliEdgeNet chọn cách thứ hai.")
r(", một tính chất mã hóa vị trí tương đối tự nhiên mà không cần học thêm tham số nào.",". Nhờ vậy, vị trí tương đối được mã hóa mà không cần thêm tham số học nào.")
r("biến thể RoPE trục (axial) cho","biến thể RoPE trục cho")
r("Chương 4 sẽ cho thấy khoản tiết kiệm nhỏ này (384 tham số trong ChiliEdgeNet) là kỹ thuật LLM duy nhất trong ba kỹ thuật phải đánh đổi bằng độ chính xác đo được, có ý nghĩa thống kê sau hiệu chỉnh Holm.",
  "Mục 4.5 kiểm tra tác động của khoản tiết kiệm nhỏ này (384 tham số trong ChiliEdgeNet) bằng ablation.")
r("thông qua ba nhánh xoay (rotate) và nén (Z-pool) song song","thông qua ba nhánh xoay tensor và nén (Z-pool) song song")
r("(kênh-cao, kênh-rộng, cao-rộng)","(kênh với chiều cao, kênh với chiều rộng, chiều cao với chiều rộng)")
r("Mục 3.2.4 trình bày việc ChiliEdgeNet loại bỏ cả hai khối này khỏi kiến trúc đề xuất, và mục 4.3 trình bày bằng chứng ablation chính thức cho quyết định đó, khác với báo cáo giai đoạn đầu của đề tài, nơi quyết định này chỉ dựa trên một khảo sát thiết kế sơ bộ.",
  "ChiliEdgeNet bỏ cả hai khối này (mục 3.2.4); bằng chứng ablation cho quyết định đó nằm ở mục 4.3.")
r(", một rủi ro đặc biệt đáng lưu ý khi phần lớn ảnh trong ChiliEdge-DS đều chụp trên nền trắng có kiểm soát.",". Rủi ro này càng đáng lưu ý vì phần lớn ảnh trong ChiliEdge-DS được chụp trên nền trắng.")
# ---- Chương 2
r("RTR_Lite_MobileNetV2 [7]","RTR_Lite_MobileNet [7]")
r("Điểm chung của cả nhóm là bản chất thuần tích chập, chỉ suy luận cục bộ tại mỗi lớp, phải xếp chồng nhiều lớp để mở rộng trường nhận thức.","Cả nhóm đều thuần tích chập, nên chịu giới hạn về trường nhận thức đã nêu ở mục 1.2.")
r("Vision Transformer khai thác trực tiếp quan hệ giữa các token, khác với CNN sâu vốn có thể mở rộng trường nhận thức để tổng hợp ngữ cảnh rộng nhưng không trực tiếp mô hình hóa quan hệ theo cặp giữa mọi vị trí; các cấu hình ViT tiêu chuẩn vì vậy thường có chi phí cao hơn.",
  "Nhóm này dùng tự chú ý để mô hình hóa quan hệ giữa các token (mục 1.3), đổi lại chi phí tính toán thường cao hơn CNN nhẹ.")
r("mười tổ hợp backbone-bộ phân loại","mười tổ hợp giữa backbone và bộ phân loại")
r("đây là một phân tích khám phá (exploratory), không nhằm khẳng định vị trí xếp hạng tuyệt đối.","đây là một phân tích khám phá và không dùng để xếp hạng.")
r("Cuối cùng, về tính minh bạch của thiết kế, các công trình thường báo cáo một mô hình đã lắp ráp hoàn chỉnh mà không đo xem từng thành phần kiến trúc có thực sự xứng đáng với số tham số nó chiếm dụng hay không, tức là không phân biệt được thành phần nào thật sự cần thiết và thành phần nào chỉ làm tăng độ phức tạp mà không cải thiện độ chính xác đo được.",
  "Cuối cùng là tính minh bạch của thiết kế. Các công trình thường báo cáo mô hình hoàn chỉnh mà không đo xem từng thành phần có xứng với số tham số nó chiếm hay không. Vì vậy khó biết thành phần nào thật sự cần, thành phần nào chỉ làm mô hình phức tạp thêm.")
# ---- Chương 3
r("dựa trên kiểu hình quan sát trực quan (phenotype-based),","dựa trên kiểu hình quan sát trực quan,")
r("CNN khai thác tốt các đặc trưng cục bộ nhưng còn hạn chế trong việc trực tiếp mô hình hóa quan hệ giữa các vùng cách xa nhau (mục 1.2), trong khi triệu chứng như xoăn lá, vàng lá lan tỏa toàn bộ bề mặt lá và cần quan hệ giữa các vùng cách xa nhau; đây đúng là thế mạnh của cơ chế tự chú ý (mục 1.3). ChiliEdgeNet vì vậy kết hợp",
  "Triệu chứng xoăn lá và vàng lá trải khắp bề mặt lá, nên mô hình cần nắm được quan hệ giữa các vùng cách xa nhau. CNN làm tốt phần đặc trưng cục bộ (mục 1.2), còn tự chú ý xử lý tốt quan hệ tầm xa (mục 1.3). ChiliEdgeNet vì thế kết hợp")
r("Việc cắt trước features[14] mang tính cơ học chứ không tùy tiện: đây là tầng làm giảm bản đồ đặc trưng xuống còn 7×7, tức số token giảm từ 196 xuống 49.",
  "Điểm cắt được đặt trước features[14] vì tầng này giảm bản đồ đặc trưng xuống 7×7, tức số token giảm từ 196 còn 49.")
r("một lớp tuyến tính (fully-connected) dự đoán","một lớp tuyến tính dự đoán")
r("Khác với báo cáo giai đoạn đầu của đề tài, nơi quyết định loại bỏ này chỉ dựa trên một khảo sát thiết kế sơ bộ, mục 4.3 (Chương 4) trình bày bằng chứng ablation chính thức: thiết kế thêm-lại theo kiểu 2×2 (cả hai khối, chỉ Triplet, chỉ RES_Block, không khối nào) kèm kiểm định McNemar có hiệu chỉnh Holm, cho quyết định loại bỏ cả hai khối này khỏi kiến trúc đề xuất.",
  "Ở giai đoạn đầu, quyết định này chỉ dựa trên một khảo sát sơ bộ. Mục 4.3 kiểm tra lại bằng thiết kế thêm lại kiểu 2×2 (cả hai khối, chỉ Triplet, chỉ RES_Block, không khối nào) kèm kiểm định McNemar có hiệu chỉnh Holm.")
r("Bảng 3.2 tổng hợp lại toàn bộ siêu tham số của khối Tiny Transformer đã đề cập rải rác ở trên, tập hợp tại đây để tiện tra cứu.","Bảng 3.2 tổng hợp siêu tham số của khối Tiny Transformer.")
# ---- Chương 4
r("Chương này trình bày thiết lập thực nghiệm và các kết quả đánh giá ChiliEdgeNet: so sánh với 17 baseline, ablation ở cấp khối và cấp từng kỹ thuật, stress test trước suy giảm ảnh, đánh giá liên bộ dữ liệu sau thích nghi few-shot, đo độ trễ trên ứng dụng camera laptop, và một khảo sát mở rộng về nén mô hình. Các so sánh chính được báo cáo kèm khoảng tin cậy Wilson và kiểm định McNemar có hiệu chỉnh Holm như đã trình bày ở mục 1.7.",
  "Mọi thí nghiệm trong chương dùng cùng một lần chia dữ liệu và seed 42. Các so sánh chính đi kèm khoảng tin cậy Wilson và kiểm định McNemar có hiệu chỉnh Holm (mục 1.7).")
r("Với mỗi phép so sánh, các bảng báo cáo hai số đếm thô, viết dưới dạng b/c, cùng giá trị p của McNemar, trong đó b là số ảnh mô hình thứ nhất dự đoán đúng còn mô hình thứ hai dự đoán sai, và c là số ảnh ngược lại.",
  "Các bảng ghi cặp số đếm b/c và giá trị p theo ký hiệu ở mục 1.7.2.")
r("Phần không phụ thuộc vào độ nhiễu của tập test là hiệu quả tham số: một con số chính xác tuyệt đối, không cần kiểm định thống kê.","Khoản giảm tham số thì không phụ thuộc vào tập test.")
r("Kết quả này ủng hộ việc giữ nguyên các nhánh skip đang hoạt động của MobileNetV2, phát hiện ablation rõ ràng nhất trong toàn bộ Chương 4.","Kết quả này ủng hộ việc giữ các kết nối tắt của MobileNetV2. Đây cũng là kết quả ablation rõ nhất trong Chương 4.")
r("toàn bộ kết nối tắt (skip connection) của các khối","toàn bộ kết nối tắt của các khối")
r(" Thích nghi chuyển miền với ít mẫu cho bài toán bệnh cây trồng cũng đã được nghiên cứu, chẳng hạn qua hiệu chỉnh thống kê theo miền đích [44] hay chuyển từ ảnh phòng thí nghiệm sang ảnh đồng ruộng [41] (mục 2.7); đánh giá ở đây xuất phát từ cùng động cơ nhưng giới hạn ở phép ánh xạ bốn lớp đặc thù cho cây ớt mô tả dưới đây.",
  " Động cơ của thí nghiệm giống các nghiên cứu ở mục 2.7.")
r("Healthy Leaf - healthy, Curl Virus - leaf_curl, Cercospora Leaf Spot - leaf_spot, Nutrition Deficiency - yellow_leaf","Healthy Leaf → healthy, Curl Virus → leaf_curl, Cercospora Leaf Spot → leaf_spot, Nutrition Deficiency → yellow_leaf")
r("Precision, recall và F1 vĩ mô (macro) chỉ tính","Precision, recall và F1 trung bình vĩ mô chỉ tính")
r("Khoảng thời gian đo bắt đầu từ lúc chuyển ảnh RGB của ROI sang định dạng PIL, tiếp tục qua resize, chuẩn hóa, suy luận mô hình, softmax và chọn lớp dự đoán top-1. Phép đo không bao gồm thời gian đọc khung hình từ camera, cắt ROI, chuyển đổi BGR sang RGB, dựng giao diện, hay ghi file.",
  "Phép đo gồm tiền xử lý, suy luận, softmax và chọn lớp top-1; giao thức chi tiết ở mục 6.4.2.")
r("Trung vị 29,70 ms và bách phân vị 95 là 46,49 ms cho thấy giai đoạn tiền xử lý cộng suy luận có thể phản hồi nhanh trên CPU laptop mà không cần GPU. Tuy nhiên, kết quả đến từ một phiên đo duy nhất trên một thiết bị duy nhất và không bao gồm toàn bộ chi phí từ lúc chụp đến lúc hiển thị của ứng dụng; thiết bị kiểm thử cũng không phải phần cứng biên chuyên dụng. Thí nghiệm này vì vậy chỉ cung cấp bằng chứng sơ bộ về tính khả thi của việc tích hợp ChiliEdgeNet vào một ứng dụng camera, không phải một phép đo đại diện cho toàn bộ pipeline hay cho phần cứng biên chuyên dụng. Dữ liệu gán nhãn trong phiên demo không cân bằng và chứa nhiều khung hình lặp lại của cùng một đối tượng, nên accuracy và ma trận nhầm lẫn không được báo cáo ở mục này.",
  "Trung vị 29,70 ms và bách phân vị 95 là 46,49 ms cho thấy phần tiền xử lý cộng suy luận phản hồi nhanh trên CPU laptop mà không cần GPU. Kết quả này chỉ đến từ một phiên đo trên một máy và chưa tính toàn bộ chi phí từ lúc chụp đến lúc hiển thị. Máy thử cũng không phải phần cứng biên chuyên dụng, nên đây mới là bằng chứng sơ bộ về khả năng tích hợp. Độ chính xác trong phiên này không được báo cáo vì dữ liệu gán nhãn không cân bằng và có nhiều khung hình lặp của cùng một lá.")
r("Mô hình đề xuất đã rất nhẹ (671.173 tham số, khoảng 2,84 MB ở FP32)","Mô hình đề xuất đã rất nhẹ (671.173 tham số, khoảng 2,7 MB trọng số ở FP32; tệp checkpoint trên đĩa 2,838 MB)")
r("vì mô hình vốn đã đủ nhỏ nên chi phí","có thể vì mô hình đã đủ nhỏ nên chi phí")
r("(PyTorch trên CPU không hỗ trợ suy luận FP16 hiệu quả cho kiến trúc này)","(trong phép thử này, PyTorch trên CPU không chạy FP16 hiệu quả cho kiến trúc này)")
# ---- Chương 5
r("kết quả này không chứng minh mô hình đã học đúng đặc trưng bệnh lý, hoàn toàn bỏ qua nền, hay nắm bắt một quan hệ nhân quả thật sự.","kết quả này chưa đủ để kết luận mô hình đã học đúng đặc trưng bệnh lý.")
r(" Lưới 14×14 cho 196 vị trí không gian, gấp bốn lần một lưới 7×7, mang lại cho bản đồ Rollout độ phân giải trực quan mịn hơn (mục 3.2.1).","")
r(", tính với 20 mẫu và hai tensor baseline hằng số (0 và 1) trong không gian đầu vào đã chuẩn hóa.",".")
r("Bản đồ được tổng hợp trên lưới 14×14 bằng cách lấy trung bình dòng chú ý trên 196 token vì mô hình dùng mean pooling, không có token CLS.","Bản đồ được tổng hợp trên lưới 14×14 theo cách mô tả ở mục 5.1.")
r("Con số xấp xỉ 105 lần đối chiếu 20,24 ms suy luận cục bộ với 2.123 ms trung vị round-trip API đám mây. Hai con số này không tương đương: round-trip bao gồm độ trễ mạng, hàng đợi phía máy chủ và gộp lô (batching) phía nhà cung cấp; không phần nào trong số đó là chi phí tính toán nội tại của mô hình. Đây là một so sánh giữa hai phương thức triển khai, không phải một khẳng định rằng một kiến trúc tính toán nhanh hơn kiến trúc kia khoảng 105 lần. Tương tự, dù độ chính xác được tính trên cùng 373 ảnh, khoảng cách độ chính xác không cấu thành một so sánh ngang hàng, vì Gemini chạy ở chế độ zero-shot còn ChiliEdgeNet được huấn luyện có giám sát riêng cho bài toán này. Endpoint thương mại cũng có thể được cập nhật theo thời gian, nên phép so sánh này chỉ có thể tái lập được đối với đúng mã định danh mô hình, prompt và ngày truy cập đã nêu.",
  "Con số khoảng 105 lần có được khi đặt 20,24 ms suy luận cục bộ cạnh 2.123 ms trung vị thời gian khứ hồi khi gọi API. Thời gian khứ hồi gồm cả độ trễ mạng, hàng đợi và cách nhà cung cấp gộp lô phía máy chủ. Vì vậy con số này phản ánh hai cách triển khai khác nhau, chưa nói lên tốc độ tính toán của hai kiến trúc. Khoảng cách độ chính xác cũng cần đọc tương tự: Gemini chạy zero-shot, còn ChiliEdgeNet được huấn luyện có giám sát cho đúng bài toán này. Endpoint thương mại có thể thay đổi theo thời gian, nên kết quả chỉ tái lập được với đúng mã mô hình, câu lệnh gợi ý và ngày truy cập đã nêu.")
# ---- Chương 6
r("Kết quả trên notebook là một chuyện, chạy được thành sản phẩm thật lại là chuyện khác. Để kiểm chứng điều đó, chuyên đề xây dựng một ứng dụng minh họa gồm backend suy luận (FastAPI, Python) phục vụ mô hình ChiliEdgeNet trực tiếp bằng FP32 và frontend (Next.js, React) hiển thị kết quả, các bản đồ diễn giải XAI (Chương 5), và một khung hỏi-đáp dùng LLM. Hai thành phần giao tiếp qua REST API cục bộ. Riêng phần tích hợp camera và đo độ trễ tiền xử lý/suy luận (mục 6.4) là phần đã được đánh giá định lượng và báo cáo chính thức trong bài báo khoa học của đề tài (mục 4.7); phần giao diện web tải ảnh, panel XAI và khung hỏi-đáp LLM trình bày dưới đây là tư liệu minh họa bổ sung, mô tả sản phẩm demo đi kèm chuyên đề, không phải một thực nghiệm được báo cáo trong bài báo.",
  "Chương này mô tả ứng dụng minh họa được xây dựng quanh ChiliEdgeNet. Ứng dụng gồm backend suy luận (FastAPI, Python) phục vụ mô hình ở định dạng FP32 và frontend (Next.js, React) hiển thị kết quả, các bản đồ diễn giải XAI (Chương 5) cùng một khung hỏi đáp dùng LLM. Hai thành phần giao tiếp qua REST API cục bộ. Trong chương này, chỉ phần tích hợp camera và đo độ trễ (mục 6.4) có số liệu định lượng, đã báo cáo ở mục 4.7. Giao diện web, khối XAI và khung hỏi đáp là phần minh họa sản phẩm, chưa được đánh giá bằng thực nghiệm.")
r(" Vì mô hình chỉ nặng 2,7 MB và không cần GPU để suy luận nhanh (mục 4.1, 4.7), backend chạy tốt trên CPU tiêu dùng."," Backend chạy được trên CPU tiêu dùng (mục 4.1, 4.7).")
r("Checkpoint được nạp một lần khi khởi động app, nên các yêu cầu về sau không phải trả thêm chi phí nạp mô hình. Vì ChiliEdgeNet chỉ nặng 2,7 MB, backend chạy tốt trên bản torch CPU mà không cần GPU.","Checkpoint được nạp một lần khi khởi động ứng dụng, nên các yêu cầu sau không mất thêm thời gian nạp mô hình.")
r("tức model còn phân vân","tức mô hình còn phân vân")
r("khẳng định chắc nịch","khẳng định chắc chắn")
r("ngay trên server trước","ngay trên máy chủ trước")
r("thư viện ảnh mẫu bấm-là-chạy","thư viện ảnh mẫu chạy ngay khi bấm")
r("Một chi tiết nhỏ nhưng cần thiết cho trải nghiệm: khi người dùng","Khi người dùng")
r("Điểm quan trọng nhất của thiết kế này là mô hình ngôn ngữ không nhìn thấy ảnh.","Mô hình ngôn ngữ không nhìn thấy ảnh.")
r("Nói cách khác, việc phân loại vẫn thuộc trọn về ChiliEdgeNet chạy cục bộ; mô hình ngôn ngữ chỉ đóng vai trò diễn đạt và không can thiệp vào quyết định.","Việc phân loại do ChiliEdgeNet đảm nhận hoàn toàn trên máy cục bộ; mô hình ngôn ngữ chỉ diễn đạt lại kết quả.")
r("Backend cài bản torch CPU: mô hình chỉ 2,7 MB nên không cần GPU, và việc bỏ toàn bộ phần CUDA giúp image nhẹ đi đáng kể.","Backend cài bản torch cho CPU; bỏ phần CUDA giúp image nhẹ đi đáng kể.")
r("Điểm đáng chú ý nhất của DFD là ranh giới dữ liệu với dịch vụ bên ngoài:","DFD cho thấy rõ ranh giới dữ liệu với dịch vụ bên ngoài:")
# ---- Chương 7
r(" Riêng khoản giảm 35% tham số không phụ thuộc vào nhiễu huấn luyện hay tiêu chí lựa chọn baseline.","")
# 544: replace remainder after bold lead (spans subscripts)
# ---- Kết luận
r("Chuyên đề này đề xuất ChiliEdgeNet, một mô hình lai CNN-Transformer 671.173 tham số, đạt 98,93% độ chính xác trên bài toán nhận diện bệnh cây ớt trong khi vẫn chạy hoàn toàn offline trên CPU. Điều đáng chú ý hơn con số tuyệt đối này là cách hai loại lợi ích so với kiến trúc nền tảng RTR_Lite_MobileNet [7] được chứng minh khác nhau: khoản tiết kiệm tham số (khoảng 35%) là một phép tính số học chính xác, không phụ thuộc vào nhiễu dữ liệu; còn mức tăng 0,80 điểm phần trăm về độ chính xác, vì mỗi kiến trúc chỉ chạy một lần và hai khoảng tin cậy chồng lấn nhau, chỉ nên được đọc như một xu hướng chứ chưa phải một xếp hạng thống kê đã được thiết lập chắc chắn. Sự phân biệt giữa cái đo được chắc chắn và cái cần diễn giải thận trọng là điều xuyên suốt cách chuyên đề trình bày kết quả ở các chương sau.",
  "Chuyên đề đề xuất ChiliEdgeNet, một mô hình lai CNN-Transformer 671.173 tham số, đạt 98,93% độ chính xác trên bài toán nhận diện bệnh cây ớt và chạy ngoại tuyến trên CPU. So với kiến trúc nền tảng RTR_Lite_MobileNet [7], hai lợi ích cần được đọc khác nhau. Khoản giảm khoảng 35% tham số là một phép đếm chính xác. Mức tăng 0,80 điểm phần trăm độ chính xác thì chỉ là xu hướng, vì mỗi mô hình chạy một lần và hai khoảng tin cậy chồng lấn nhau.")
r("Đóng góp phương pháp trọng tâm là kiểm toán thiết kế theo từng thành phần thay vì chỉ báo cáo một mô hình đã lắp ráp hoàn chỉnh, bằng ablation huấn luyện lại từ đầu, kiểm định ghép cặp từng ảnh có hiệu chỉnh Holm, và một đánh giá chuyển miền trên bộ dữ liệu độc lập. Bốn kết quả nổi bật xuyên suốt Chương 4 và Chương 5:",
  "Về phương pháp, chuyên đề kiểm tra từng thành phần thiết kế bằng ablation huấn luyện lại từ đầu, kiểm định ghép cặp có hiệu chỉnh Holm và một đánh giá chuyển miền trên bộ dữ liệu độc lập. Có bốn kết quả chính:")
r("1. Các kết nối tắt của MobileNetV2 quan trọng hơn khối Tiny Transformer mới thêm vào.","1. Trên ảnh sạch, kết nối tắt của MobileNetV2 có tác động lớn hơn khối Tiny Transformer.")
r(", gợi ý vai trò của khối này là duy trì biểu diễn dưới suy giảm mạnh hơn là mang lại lợi thế ổn định trên ảnh sạch.",". Khối này có vẻ giúp giữ biểu diễn khi ảnh suy giảm mạnh hơn là tăng độ chính xác trên ảnh sạch.")
r("2. Ba kỹ thuật tiết kiệm tham số từ LLM chưa cho thấy phải đánh đổi độ chính xác trên ảnh sạch, ở dữ liệu một lần chạy này.","2. Ở lần chạy này, ba kỹ thuật tiết kiệm tham số từ LLM không làm giảm độ chính xác trên ảnh sạch.")
r("3. Không phải khối kế thừa nào cũng có ích.","3. Hai khối kế thừa Triplet Attention và RES_Block không mang lại lợi ích đo được.")
r("4. Thích nghi miền hoạt động rất hiệu quả với chỉ một tập hỗ trợ nhỏ.","4. Một tập hỗ trợ nhỏ đủ để cải thiện mạnh kết quả chuyển miền.")
r(", một cải thiện lớn hơn một bậc độ lớn so với bất kỳ thành phần kiến trúc nào đo được ở Chương 4, dù kết quả này chỉ áp dụng cho bốn lớp được ánh xạ và một bộ dữ liệu ngoài miền duy nhất.",". Mức cải thiện này lớn hơn một bậc so với tác động của bất kỳ thành phần kiến trúc nào ở Chương 4, dù chỉ áp dụng cho bốn lớp được ánh xạ và một bộ dữ liệu ngoài miền.")
r("Điểm cắt xương sống vẫn là một quyết định kiến trúc quan trọng: giữ lại features[0:13]","@@P562@@")

for old,new,c in R:
    if new=="@@P562@@": continue
    rep(old,new,c)

# P562 full rewrite
for p in d.paragraphs:
    if p.text.startswith("Điểm cắt xương sống vẫn là một quyết định kiến trúc quan trọng"):
        new=("Điểm cắt xương sống tại features[0:14] giữ được lưới đặc trưng 14×14 (196 token), nhờ đó bản đồ Attention Rollout ở Chương 5 đủ mịn để khoanh vùng triệu chứng. "
             "Trên năm ảnh minh họa, vùng phản hồi cao của ba kỹ thuật XAI nhìn chung trùng với lá, quả hoặc vùng bệnh. "
             "Trên toàn bộ 373 ảnh test, ChiliEdgeNet đạt 98,93% so với 78,55% của Gemini 3.1 Flash Lite chạy zero-shot. "
             "Độ trễ ở mức hệ thống thấp hơn khoảng 105 lần, dù đây là so sánh giữa hai cách triển khai (mục 5.2). "
             "Mục 4.7 và Chương 6 cho thấy mô hình tích hợp được vào một ứng dụng camera, với độ trễ trung vị dưới 30 ms trên CPU laptop. "
             "Các hướng phát triển ở mục 7.2 gồm xác nhận kết quả trên nhiều seed, thử với suy giảm thu nhận ảnh thật, đo trên phần cứng biên vật lý và mở rộng đánh giá chuyển miền.")
        para_replace(p,p.text,new); log.append("P562 rewritten")
# P544 remainder
for p in d.paragraphs:
    if p.text.startswith("Số lần chạy và năng lực thống kê."):
        full=''.join(r.text for r in p.runs)
        k=full.find("Mô hình đề xuất và toàn bộ ablation")
        para_replace(p, full[k:], "Mô hình đề xuất và toàn bộ ablation đều đến từ một lần chạy seed 42. Chỉ hai kết quả còn ý nghĩa sau hiệu chỉnh Holm: tác động của kết nối tắt trên ảnh sạch và ưu thế của khối Tiny Transformer ở mức giảm sáng 0,2 (mục 4.3, 4.4). Các chênh lệch còn lại, kể cả ba phép so sánh Triplet/RES_Block, mới dừng ở mức xu hướng và cần được xác nhận qua nhiều lần chạy.")
        log.append("P544 rewritten")
# delete contribution bullet 6
for p in d.paragraphs:
    if "Hệ thống hóa cơ sở lý thuyết đầy đủ cho toàn bộ kỹ thuật" in p.text:
        p._p.getparent().remove(p._p); log.append("deleted bullet 'Hệ thống hóa'")

# ---- global replacements in body text (before references)
paras=d.paragraphs
ref_i=next(i for i,p in enumerate(paras) if p.text.strip().upper().startswith("TÀI LIỆU THAM KHẢO") and i>300)
BODY=set(id(p._p) for p in paras[:ref_i])
inbody=lambda p: id(p._p) in BODY
for old,new in [("hỏi-đáp","hỏi đáp"),("thêm-lại","thêm lại"),("kéo-thả","kéo thả"),("trong-miền","trong miền"),("H=4, G=2","H = 4, G = 2"),
                ("con số accuracy","giá trị độ chính xác"),("offline","ngoại tuyến")]:
    h=rep(old,new,0,inbody); log.append(f"{old} -> {new}: {h}")
# lowercase accuracy word -> độ chính xác
n=0
for p in paras[:ref_i]:
    for r_ in p.runs:
        if re.search(r'(?<![A-Za-z_\-])accuracy(?![A-Za-z_\-])',r_.text):
            r_.text=re.sub(r'(?<![A-Za-z_\-])accuracy(?![A-Za-z_\-])','độ chính xác',r_.text); n+=1
log.append(f"accuracy runs: {n}")
# ranges digit-digit -> en dash
n=0; ex=[]
for p in paras[:ref_i]:
    for r_ in p.runs:
        t=r_.text
        new=re.sub(r'(?<![A-Za-z])(\d[\d.,]*)-(?=\d)',lambda m:m.group(1)+'–',t)
        if new!=t: ex.append(t[max(0,t.find('-')-15):t.find('-')+15]); r_.text=new; n+=1
log.append(f"range dashes: {n} :: {ex}")
d.save(DST)
print("\n".join(log))
